"""Build the meeting deck as a PowerPoint file on the lab template.

    python make_pptx.py        (from meeting_feedback/presentation)

Writes Interlocking_by_curvature_meeting.pptx next to this file, using the
master of ../ppt_tmpl_v2025.2 - Copy - Copy - Copy.pptx (gradient bar, topic
and slide number). The numbers come from pairs.json (the per-R DBG results).
"""
import copy
import json
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "ppt_tmpl_v2025.2 - Copy - Copy - Copy.pptx")
import sys
# Optional first argument: another output name (e.g. when the deck is open in PowerPoint).
OUT = os.path.join(HERE, sys.argv[1] if len(sys.argv) > 1 else "Interlocking_by_curvature_meeting.pptx")
DATA = json.load(open(os.path.join(HERE, "pairs.json"), encoding="utf-8"))

BLUE = RGBColor(0x0C, 0x7C, 0xAE)
LIGHT = RGBColor(0xDC, 0xEB, 0xF5)
GREY = RGBColor(0x5B, 0x63, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1D, 0x1D, 0x1F)

prs = Presentation(TEMPLATE)

# Drop the template's own example slides; keep its master and layouts.
for sldId in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sldId.rId)
    prs.slides._sldIdLst.remove(sldId)

L_TITLE = prs.slide_layouts[0]
L_CONTENT = prs.slide_layouts[1]

# Content area of the "title and content" layout.
X0, Y0, W, H = Inches(0.37), Inches(1.24), Inches(12.6), Inches(5.65)


def add_slide(topic, title, notes=""):
    s = prs.slides.add_slide(L_CONTENT)
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text = title
        elif idx == 13:
            ph.text = topic
    # The layout's slide number is not copied by add_slide; bring it over.
    for ph in L_CONTENT.placeholders:
        if ph.placeholder_format.idx == 12:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def body(s):
    return next(ph for ph in s.placeholders if ph.placeholder_format.idx == 1)


def drop_body(s):
    b = body(s)
    b._element.getparent().remove(b._element)


def bullets(s, items, size=18, box=None):
    """items: list of (text, level) or str; '**x**' makes x bold."""
    if box is None:
        tf = body(s).text_frame
    else:
        drop_body(s)
        tf = s.shapes.add_textbox(*box).text_frame
        tf.word_wrap = True
    tf.clear()
    first = True
    for it in items:
        text, level = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        for k, part in enumerate(re.split(r"\*\*", text)):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size - 2 * level)
            r.font.bold = (k % 2 == 1)
        p.space_after = Pt(6)
    return tf


def textbox(s, x, y, w, h, text, size=14, bold=False, color=None, align=None):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for k, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = tf.paragraphs[0].add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.bold = bold or (k % 2 == 1)
        if color is not None:
            r.font.color.rgb = color
    if align is not None:
        tf.paragraphs[0].alignment = align
    return tb


def table(s, x, y, w, head, rows, widths, size=12, row_h=Inches(0.36)):
    n, m = len(rows) + 1, len(head)
    gt = s.shapes.add_table(n, m, x, y, w, row_h * n).table
    total = sum(widths)
    for j, wj in enumerate(widths):
        gt.columns[j].width = Emu(int(w * wj / total))
    for j, h in enumerate(head):
        c = gt.cell(0, j)
        c.text = h
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            c = gt.cell(i, j)
            c.text = str(v)
            c.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    for i in range(n):
        gt.rows[i].height = row_h
    gt.rows[0].height = min(row_h, Inches(0.45))
    return gt


def placeholder_box(s, x, y, w, h, label):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    b.fill.background()
    b.line.color.rgb = RGBColor(0x9A, 0xA6, 0xB2)
    b.line.dash_style = MSO_LINE.DASH
    tf = b.text_frame
    tf.text = label
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].runs[0].font.size = Pt(14)
    tf.paragraphs[0].runs[0].font.color.rgb = GREY
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


# ------------------------------------------------------------------ data
def fmtR(R):
    return "%g" % R


def blk(r):
    b = ["%d–%d" % (p["a"], p["b"]) for p in r["pairs"] if p["blocked"]]
    return ", ".join(b) if b else "none"


def key_row(r):
    if r["detJ"] <= 0:
        return [fmtR(r["R"]), "%.3f folds" % r["detJ"], "–", "–", "invalid"]
    imm = sum(1 for c in r["pclr"] if c < 0)
    return [fmtR(r["R"]), "%.3f" % r["detJ"], blk(r), "%d of 8" % imm, "no"]


def crit_row(r):
    if r["detJ"] <= 0:
        return [fmtR(r["R"]), "no (folds)", "–", "–", "–"]
    t = r["thin"]
    q = r["ratio"]
    return [fmtR(r["R"]), "yes", "%.1f%% %s" % (t, "pass" if t >= 3 else "fail"),
            "%.1f %s" % (q, "pass" if q <= 2 else ("warn" if q <= 3 else "fail")), "no / 0"]


SPH = [r for r in DATA if r["shape"] == "sphere"]
DUCK = [r for r in DATA if r["shape"] == "duck"]

# ------------------------------------------------------------------ 1 title
s = prs.slides.add_slide(L_TITLE)
s.shapes.title.text = "Interlocking by curvature"
sub = next(ph for ph in s.placeholders if ph.placeholder_format.idx == 1)
sub.text_frame.text = "그래픽스 및 기하처리 연구실"
p = sub.text_frame.add_paragraph()
p.text = "정보컴퓨터공학부 예기즈바에브 졸란"

# ------------------------------------------------------------------ 2 problem statement
STATEMENT = (
    "Given a closed genus-0 solid, partition it into k pieces that interlock in R³: "
    "no proper subassembly separates by translation except one designated key.",
    "Elber's construction divides the parameter domain D of a trivariate M and maps each cell "
    "through M. If M is affine the cuts stay planar, and a planar partition always comes apart, "
    "so joints are the only resistance. We cut in D and let M curve the cuts. Two pieces sharing "
    "an interface S separate along d only if d·n ≥ 0 at every point of S; when the normals spread "
    "far enough, no such d exists and the pair is locked.",
    "**Criteria:** interlocking verified in R³ (single key, then level k); Jacobian of M∘R "
    "positive everywhere; every piece one connected component above a minimum thickness; "
    "roughly uniform piece sizes; no joints exposed in cavities.",
)
words = len(re.findall(r"[A-Za-z0-9R³∘·≥'-]+", " ".join(STATEMENT).replace("**", "")))
assert words <= 150, words
s = add_slide("Problem statement", "Problem statement (%d words)" % words,
              "Scope for the first round: genus 0, no limbs, translation only.")
bullets(s, list(STATEMENT), size=18)
for p in body(s).text_frame.paragraphs:
    p.level = 0

# ------------------------------------------------------------------ 3 setup (slide 5)
s = add_slide("Step 1 · Reproduce Elber, pins removed", "Two models from Elber's own script",
              "Only these two models are used. Elber's grids are N×M×1; the professor asked for "
              "2×2×2, which adds the radial cut w through the thickness.")
bullets(s, [
    "Elber's **puz_vol.irt** (IRIT), copied as puz_vol - Copy.irt with every pin and hole deleted",
    "**PuzzleVolEarth**: a sphere shell.  **PuzzleVolDuck**: ruled between the duck surface and an inner core",
    "Divided **2×2×2** in D: u around the axis, v along the profile, w through the thickness",
    "Tiles at size 1.0 (Elber's 0.99 left room for pins); the trivariate M is saved next to the pieces",
], size=18, box=(X0, Y0, Inches(6.2), H))
placeholder_box(s, Inches(6.8), Y0 + Inches(0.1), Inches(6.1), Inches(5.3),
                "[Screenshot: GuIrit, PuzzleVolEarth and PuzzleVolDuck, pins removed]")

# ------------------------------------------------------------------ 4 R sweep (slide 6)
s = add_slide("Step 3 · The R sweep", "R is a randomization amount, not a radius",
              "Re-seeding makes the sweep one shape at several strengths. Folded cases are kept to "
              "show what folding does.")
bullets(s, [
    "RandomizeInterior(TV, R): every **interior** control point += random(−R, R); the outer shape is unchanged",
    "Random generator re-seeded before every R: each R scales **one fixed** field of offsets",
    "Duck refined in depth first (it had only 2 control points across its thickness)",
    "Rule: interior **det J ≤ 0** → the map folds → the case is invalid and not ranked",
], size=18, box=(X0, Y0, Inches(7.0), H))
table(s, Inches(7.4), Y0 + Inches(0.2), Inches(5.5), ["model", "valid R", "folds from"],
      [["sphere", "0 – 0.125", "0.15"], ["duck", "0 – 0.015", "0.02"],
       ["Elber's Earth, R = 0.2", "–", "folds"]], [48, 26, 26], size=14, row_h=Inches(0.45))

# ------------------------------------------------------------------ 5 the table
s = add_slide("Step 4 · The table", "Blocking pairs, immobilized pieces, single key",
              "Blocked or free is decided exactly by the clearance (distance from the origin to the "
              "convex hull of the face normals), not by sampling directions.")
drop_body(s)
head = ["R", "det J", "blocking pairs", "immobilized", "single key"]
wid = [12, 20, 32, 20, 16]
textbox(s, X0, Y0, Inches(6.1), Inches(0.4), "Sphere (PuzzleVolEarth)", size=16, bold=True)
table(s, X0, Y0 + Inches(0.42), Inches(6.1), head, [key_row(r) for r in SPH], wid, size=12)
textbox(s, Inches(6.8), Y0, Inches(6.1), Inches(0.4), "Duck (PuzzleVolDuck)", size=16, bold=True)
table(s, Inches(6.8), Y0 + Inches(0.42), Inches(6.1), head, [key_row(r) for r in DUCK], wid, size=12)
textbox(s, X0, Inches(5.45), Inches(12.5), Inches(1.5),
        "**No valid case has a single key.**  Duck: every piece is immobilized, but each piece and its "
        "radial partner slide out together.  Sphere: R = 0 blocks only on a knife edge; at R = 0.025–0.05 "
        "all four outer shells come out alone (under 2° of clearance); from R = 0.075 pair 6–7 closes.\n"
        "Pairs 0–1, 2–3, 4–5, 6–7 = the two layers of one quarter (the radial cut).  Controls: flat box "
        "0 of 12 blocked; duck 2×2×1 (Elber's topology) 0 of 4 blocked.", size=13)

# ------------------------------------------------------------------ 6 who blocks whom
s = add_slide("Step 4 · The blocking graph", "Who blocks whom",
              "Each piece touches three others. On the duck the radial neighbour always blocks, so "
              "every piece is immobilized, yet the pair moves out together through the flat cuts.")
drop_body(s)


def matrix(s, x, y, r, label):
    blocked = {(p["a"], p["b"]) for p in r["pairs"] if p["blocked"]}
    blocked |= {(b, a) for a, b in blocked}
    touch = {(p["a"], p["b"]) for p in r["pairs"]}
    touch |= {(b, a) for a, b in touch}
    textbox(s, x, y, Inches(4.6), Inches(0.4), label, size=16, bold=True)
    cell = Inches(0.46)
    gt = s.shapes.add_table(9, 9, x, y + Inches(0.45), cell * 9, cell * 9).table
    for j in range(9):
        gt.columns[j].width = cell
        gt.rows[j].height = cell
    for i in range(9):
        for j in range(9):
            c = gt.cell(i, j)
            if i == 0 and j == 0:
                c.text = ""
            elif i == 0 or j == 0:
                c.text = str((j if i == 0 else i) - 1)
            else:
                a, b = i - 1, j - 1
                if (a, b) in blocked:
                    c.text, fill, col = "B", BLUE, WHITE
                elif (a, b) in touch:
                    c.text, fill, col = "·", LIGHT, INK
                else:
                    c.text, fill, col = "", WHITE, INK
                c.fill.solid()
                c.fill.fore_color.rgb = fill
                if c.text:
                    c.text_frame.paragraphs[0].runs[0].font.color.rgb = col
            if c.text:
                c.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
                c.text_frame.paragraphs[0].runs[0].font.bold = True
            c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    imm = [str(k) for k, c in enumerate(r["pclr"]) if c < 0]
    free = [str(k) for k, c in enumerate(r["pclr"]) if c >= 0]
    textbox(s, x, y + Inches(0.55) + cell * 9, Inches(4.8), Inches(0.5),
            "Immobilized: %s.  Free alone: %s." % (", ".join(imm) or "none", ", ".join(free) or "none"), size=13)


matrix(s, Inches(0.9), Y0, next(r for r in DATA if r["file"] == "tvs_sphere_r010"), "Sphere, R = 0.1")
matrix(s, Inches(7.2), Y0, next(r for r in DATA if r["file"] == "tvs_duck_r000"), "Duck, R = 0")
textbox(s, X0, Inches(6.55), Inches(12.5), Inches(0.5),
        "B = blocked (no direction separates them)   · = touching, can separate   empty = not touching.   "
        "Piece = 4·(u half) + 2·(v half) + (w layer)", size=12, color=GREY)

# ------------------------------------------------------------------ 7 criteria per R
s = add_slide("Criteria", "The criteria, per R",
              "The 3% thickness limit is a default, not a printing requirement.")
drop_body(s)
head = ["R", "det J positive", "thinnest (3% limit)", "largest / smallest", "key / level k"]
wid = [13, 18, 25, 24, 20]
textbox(s, X0, Y0, Inches(6.1), Inches(0.4), "Sphere", size=16, bold=True)
table(s, X0, Y0 + Inches(0.42), Inches(6.1), head, [crit_row(r) for r in SPH], wid, size=12)
textbox(s, Inches(6.8), Y0, Inches(6.1), Inches(0.4), "Duck", size=16, bold=True)
table(s, Inches(6.8), Y0 + Inches(0.42), Inches(6.1), head, [crit_row(r) for r in DUCK], wid, size=12)
textbox(s, X0, Inches(5.45), Inches(12.5), Inches(1.2),
        "**One connected component:** yes in every valid row (a piece is one box through a continuous map; "
        "only a fold could break it).   **Joints in cavities:** not applicable, no joints in this experiment.   "
        "Sizes: pass ≤ 2, warn ≤ 3.", size=13)

# ------------------------------------------------------------------ 8 Elber section 6
s = add_slide("Elber 2025, §6 Future Work", "Each §6 item, matched to our plan",
              "Item 2 is where this work sits: Elber ensures (dis)assemblability manually; the DBG automates "
              "the check, so far for uniform grids and translations. The duck's radial pairs are the case "
              "he warns about: a puzzle that cannot be (dis)assembled. Items 3 and 4 need joints, which this "
              "experiment deliberately removed.")
drop_body(s)
table(s, X0, Y0, Inches(12.5), ['Elber §6 says', 'our plan', 'status now'], [['1 · Other domain topologies: hexagonal prisms, semi-regular or general space-filling cells, with matching joints', "Not in our plan. We keep Elber's cuboid grid; we change whether the cells interlock, not their shape", '–'], ['2 · Non-regular division, plus automatic verification of a collision-free (dis)assembly ("herein we manually ensure that")', 'Our core step: the DBG tests every group of pieces for a free translation and builds a take-apart sequence', 'Done for uniform 2×2×2 (complete up to 18 pieces), translation only. Not yet: non-regular division, rotations'], ['3 · Better joints than pins and holes; deep cavities can expose joints', 'Later, only if curvature alone does not lock: the DBG shows which faces are loose, only those get a joint', 'Not started: no joints in this experiment; dovetails exist only in the older planar pipeline; no cavity check'], ["4 · Pin and hole sizes distorted by M; compensate with M's derivatives", "Later, together with joints: the analyser already evaluates M's derivatives (J) at any point", 'Not started: no compensation implemented']], [36, 33, 31], size=14, row_h=Inches(1.0))

# ------------------------------------------------------------------ 9 open questions
s = add_slide("Open questions", "Open questions")
bullets(s, [
    "**Rotations?** The d·n test covers straight pulls only. Proposal: translation only for the first "
    "paper, stated as a limitation, plus a swept-volume check on the final design",
    "**Which shapes?** Proposal: genus 0, no limbs (deformed sphere, simple busts; not the torus or armadillo). "
    "Is \"no limbs\" about genus, the medial axis, or convex cross-sections? It needs an operational test",
    "**Is random R the right instrument?** It folds the duck at 0.02 and never makes the u / v cuts block. "
    "Design the curvature instead (wavy cuts in D, harmonic trivariate)?",
    "**Which level k?** Song 2012 (unique key at each step) is computed; Chen 2022 (moves before the "
    "first piece comes free) needs motion search",
    "**What do joints become?** If curvature does not lock, the DBG says which faces are loose: fewest "
    "joints, or least visible?",
    "**Fitted or constructed trivariates, and a thickness limit in mm?**",
], size=16)

# The template's Korean text rules let a Latin word break anywhere ("ran dom").
# Switch that off on every paragraph of every slide.
from pptx.oxml.ns import qn
for sl in prs.slides:
    for p in sl.shapes._spTree.iter(qn("a:p")):
        pPr = p.find(qn("a:pPr"))
        if pPr is None:
            pPr = p.makeelement(qn("a:pPr"), {})
            p.insert(0, pPr)
        pPr.set("latinLnBrk", "0")
        pPr.set("eaLnBrk", "1")

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides), "slides, statement", words, "words")
